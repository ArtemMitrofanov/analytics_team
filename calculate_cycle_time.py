#!/usr/bin/env python3
"""
Cycle Time Calculation Script
Calculates working hours between status changes for tasks,
excluding weekends, Russian holidays, and lunch break (13:00-14:00).
Working hours: 9:00-18:00 (8 working hours/day after lunch deduction).

Outputs:
1. 'Итоги' sheet - per task total hours/days
2. 'цикл тайм тестирование' sheet - overall statistics (mean, median, percentiles)
3. 'цикл тайм по периодам' sheet - statistics by month and quarter
"""

import openpyxl
from openpyxl.styles import numbers
from datetime import datetime, time, timedelta
import holidays
import statistics
from collections import defaultdict


# ==================== CONFIGURATION ====================
WORK_START = time(9, 0)
WORK_END = time(18, 0)
LUNCH_START = time(13, 0)
LUNCH_END = time(14, 0)
WORK_HOURS_PER_DAY = 8  # 9 hours - 1 hour lunch

# Russian holidays
RU_HOLIDAYS = holidays.RU()

# Input/Output file
INPUT_FILE = "YT._Задачи_по_статусам_тестирования_(без_привязки_к_Юзеру)_2026_08_26.xlsx"
SOURCE_SHEET = "result"


# ==================== HELPER FUNCTIONS ====================
def is_workday(dt: datetime) -> bool:
    """Check if date is a working day (not weekend, not holiday)."""
    if dt.weekday() >= 5:  # Saturday=5, Sunday=6
        return False
    if dt.date() in RU_HOLIDAYS:
        return False
    return True


def calculate_work_hours(start_dt: datetime, end_dt: datetime) -> float:
    """
    Calculate working hours between two datetimes.
    Excludes weekends, holidays, non-working hours, and lunch break.
    """
    if start_dt >= end_dt:
        return 0.0

    total_hours = 0.0
    current = start_dt

    while current < end_dt:
        # Skip non-working days
        if not is_workday(current):
            current = datetime.combine(current.date() + timedelta(days=1), WORK_START)
            continue

        day_start = datetime.combine(current.date(), WORK_START)
        day_end = datetime.combine(current.date(), WORK_END)

        # Adjust current to work start if before
        if current < day_start:
            current = day_start

        # Skip if already past work end
        if current >= day_end:
            current = datetime.combine(current.date() + timedelta(days=1), WORK_START)
            continue

        # Calculate work hours for this day
        day_end_actual = min(end_dt, day_end)
        if day_end_actual > current:
            day_hours = (day_end_actual - current).total_seconds() / 3600

            # Deduct lunch break if work period overlaps with 13:00-14:00
            lunch_start_dt = datetime.combine(current.date(), LUNCH_START)
            lunch_end_dt = datetime.combine(current.date(), LUNCH_END)

            if current < lunch_end_dt and day_end_actual > lunch_start_dt:
                overlap_start = max(current, lunch_start_dt)
                overlap_end = min(day_end_actual, lunch_end_dt)
                if overlap_end > overlap_start:
                    lunch_overlap = (overlap_end - overlap_start).total_seconds() / 3600
                    day_hours -= lunch_overlap

            total_hours += day_hours

        # Move to next day 9:00
        current = datetime.combine(current.date() + timedelta(days=1), WORK_START)

    return total_hours


def parse_date(value):
    """Parse date from string or return datetime if already parsed."""
    if isinstance(value, str):
        return datetime.strptime(value, '%Y-%m-%d %H:%M:%S')
    return value


# ==================== MAIN CALCULATION ====================
def main():
    print(f"Loading workbook: {INPUT_FILE}")
    wb = openpyxl.load_workbook(INPUT_FILE)
    ws = wb[SOURCE_SHEET]

    # ---- Step 1: Calculate hours per task per period from raw data ----
    task_period_hours = defaultdict(lambda: defaultdict(float))  # task -> period -> hours
    task_total_hours = defaultdict(float)

    print("Processing raw data...")
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, values_only=True):
        issue_id, stage, end_date, begin_date = row
        if not end_date or not begin_date:
            continue

        end_date = parse_date(end_date)
        begin_date = parse_date(begin_date)

        hours = calculate_work_hours(begin_date, end_date)
        task_total_hours[issue_id] += hours

        # Attribute to period by BEGIN date (when work started)
        period_month = begin_date.strftime('%Y-%m')
        period_quarter = f"{begin_date.year}-Q{(begin_date.month - 1) // 3 + 1}"

        task_period_hours[issue_id][period_month] += hours
        task_period_hours[issue_id][period_quarter] += hours

    print(f"Processed {len(task_total_hours)} tasks")

    # ---- Step 2: Create 'Итоги' sheet ----
    print("Creating 'Итоги' sheet...")
    if 'Итоги' in wb.sheetnames:
        wb.remove(wb['Итоги'])
    ws_summary = wb.create_sheet('Итоги')

    ws_summary.append(['Задача', 'Затраченое время (часы)', 'Затраченое время (дни)'])

    for task, hours in sorted(task_total_hours.items()):
        days = hours / WORK_HOURS_PER_DAY
        ws_summary.append([task, round(hours, 2), round(days, 2)])

    # Format column B as [h]:mm:ss
    for row in ws_summary.iter_rows(min_row=2, max_row=ws_summary.max_row, min_col=2, max_col=2):
        for cell in row:
            if cell.value is not None:
                cell.value = cell.value / 24  # Convert hours to Excel time (fraction of day)
                cell.number_format = '[h]:mm:ss'

    # ---- Step 3: Create 'цикл тайм тестирование' sheet ----
    print("Creating 'цикл тайм тестирование' sheet...")
    if 'цикл тайм тестирование' in wb.sheetnames:
        wb.remove(wb['цикл тайм тестирование'])
    ws_stats = wb.create_sheet('цикл тайм тестирование')

    hours_list = list(task_total_hours.values())
    days_list = [h / WORK_HOURS_PER_DAY for h in hours_list]

    ws_stats.append(['Метрика', 'Значение (часы)', 'Значение (дни)'])
    ws_stats.append(['Количество задач', len(hours_list), ''])
    ws_stats.append(['Среднее (Mean)', round(statistics.mean(hours_list), 2), round(statistics.mean(days_list), 2)])
    ws_stats.append(['Медиана', round(statistics.median(hours_list), 2), round(statistics.median(days_list), 2)])
    ws_stats.append(['Минимум', round(min(hours_list), 2), round(min(days_list), 2)])
    ws_stats.append(['Максимум', round(max(hours_list), 2), round(max(days_list), 2)])

    if len(hours_list) > 1:
        ws_stats.append(['Стандартное отклонение', round(statistics.stdev(hours_list), 2), round(statistics.stdev(days_list), 2)])
    else:
        ws_stats.append(['Стандартное отклонение', 0, 0])

    # Percentiles
    sorted_hours = sorted(hours_list)
    sorted_days = sorted(days_list)
    for p in [25, 50, 75, 90, 95]:
        idx = int(len(sorted_hours) * p / 100)
        if idx >= len(sorted_hours):
            idx = len(sorted_hours) - 1
        ws_stats.append([f'Перцентиль {p}%', round(sorted_hours[idx], 2), round(sorted_days[idx], 2)])

    # ---- Step 4: Create 'цикл тайм по периодам' sheet ----
    print("Creating 'цикл тайм по периодам' sheet...")
    if 'цикл тайм по периодам' in wb.sheetnames:
        wb.remove(wb['цикл тайм по периодам'])
    ws_period = wb.create_sheet('цикл тайм по периодам')

    # Aggregate by month and quarter
    month_data = defaultdict(list)
    quarter_data = defaultdict(list)

    for task, periods in task_period_hours.items():
        for period, hours in periods.items():
            if '-Q' in period:
                quarter_data[period].append(hours)
            else:
                month_data[period].append(hours)

    # Header
    headers = ['Период', 'Тип', 'Задач', 'Среднее (ч)', 'Медиана (ч)', 'Мин (ч)', 'Макс (ч)', 'Среднее (дни)', 'Медиана (дни)']

    # Monthly stats
    ws_period.append(headers)
    for period in sorted(month_data.keys()):
        vals = month_data[period]
        days_vals = [v / WORK_HOURS_PER_DAY for v in vals]
        ws_period.append([
            period, 'Месяц', len(vals),
            round(statistics.mean(vals), 2), round(statistics.median(vals), 2),
            round(min(vals), 2), round(max(vals), 2),
            round(statistics.mean(days_vals), 2), round(statistics.median(days_vals), 2)
        ])

    # Quarterly stats
    ws_period.append([])  # blank row
    ws_period.append(headers)
    for period in sorted(quarter_data.keys()):
        vals = quarter_data[period]
        days_vals = [v / WORK_HOURS_PER_DAY for v in vals]
        ws_period.append([
            period, 'Квартал', len(vals),
            round(statistics.mean(vals), 2), round(statistics.median(vals), 2),
            round(min(vals), 2), round(max(vals), 2),
            round(statistics.mean(days_vals), 2), round(statistics.median(days_vals), 2)
        ])

    # ---- Save ----
    wb.save(INPUT_FILE)
    print(f"\nDone! Saved to {INPUT_FILE}")
    print(f"Total tasks: {len(task_total_hours)}")
    print(f"Months: {len(month_data)}, Quarters: {len(quarter_data)}")


if __name__ == '__main__':
    main()